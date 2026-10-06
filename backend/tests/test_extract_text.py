import io

import docx
import pytest
from fpdf import FPDF

from app.profile.extract_text import ExtractError, extract_text

RESUME = (
    "Jane Doe\njane@example.com\n\nExperience\nSenior Engineer, Acme Corp  2020 - Present\n"
    "Built a billing API that cut invoice errors by 40%.\n"
)


def make_pdf(text: str) -> bytes:
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=12)
    if text:
        pdf.multi_cell(0, 8, text)
    return bytes(pdf.output())


def make_docx(*, header: str = "", table: list[list[str]] | None = None) -> bytes:
    document = docx.Document()
    for line in RESUME.splitlines():
        document.add_paragraph(line)
    if header:
        document.sections[0].header.paragraphs[0].text = header
    if table:
        grid = document.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, value in enumerate(row):
                grid.cell(r, c).text = value
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_pdf_text_is_extracted():
    result = extract_text("resume.pdf", make_pdf(RESUME))
    assert "Senior Engineer" in result.text
    assert "cut invoice errors by 40%" in result.text


def test_image_only_pdf_is_rejected_with_ats_explanation():
    with pytest.raises(ExtractError, match="scanned or image-only"):
        extract_text("scan.pdf", make_pdf(""))


def test_corrupt_pdf_is_rejected():
    with pytest.raises(ExtractError):
        extract_text("broken.pdf", b"%PDF-1.4 this is not a real pdf")


def test_docx_paragraphs_are_extracted():
    result = extract_text("resume.docx", make_docx())
    assert "Senior Engineer, Acme Corp" in result.text
    assert not any("header" in w or "Tables" in w for w in result.warnings)


def test_docx_header_text_is_kept_and_flagged():
    result = extract_text("resume.docx", make_docx(header="Jane Doe | 555-0100"))
    assert "555-0100" in result.text
    assert any("header or footer" in w for w in result.warnings)


def test_docx_tables_are_read_and_flagged():
    result = extract_text("resume.docx", make_docx(table=[["Python", "Expert"], ["Azure", "Advanced"]]))
    assert "Python | Expert" in result.text
    assert any("Tables" in w for w in result.warnings)


def test_plain_text_is_accepted():
    assert "Jane Doe" in extract_text("resume.txt", RESUME.encode()).text


def test_old_doc_and_unknown_types_are_rejected():
    with pytest.raises(ExtractError, match=".docx"):
        extract_text("resume.doc", b"x" * 100)
    with pytest.raises(ExtractError, match="Unsupported"):
        extract_text("resume.png", b"x" * 100)


def test_thin_text_gets_a_warning():
    result = extract_text("short.txt", b"Jane Doe, engineer with ten years of experience building APIs.")
    assert any("Very little text" in w for w in result.warnings)

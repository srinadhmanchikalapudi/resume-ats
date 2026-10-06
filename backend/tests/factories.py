"""Builders for the sample resume files used across the tests."""

import io

import docx
from fpdf import FPDF

RESUME = (
    "Jane Doe\njane@example.com\n\nExperience\nSenior Engineer, Acme Corp  2020 - Present\n"
    "Built a billing API that cut invoice errors by 40%.\n"
)


def make_pdf(text: str) -> bytes:
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=12)
    if text:
        pdf.multi_cell(0, 8, text)
    return bytes(pdf.output())


def make_docx(*, header: str = "", table: list[list[str]] | None = None) -> bytes:
    document = docx.Document()
    for line in RESUME.splitlines():
        document.add_paragraph(line)
    if header:
        document.sections[0].header.paragraphs[0].text = header
    if table:
        grid = document.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, value in enumerate(row):
                grid.cell(r, c).text = value
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()

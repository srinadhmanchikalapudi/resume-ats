import io

import docx
import pytest
from factories import make_docx
from fpdf import FPDF

from app import llm
from app.export.blocks import build_blocks
from app.export.docx_builder import build_docx
from app.export.pdf_builder import build_pdf
from app.profile.file_checks import check_docx, check_pdf, check_upload
from app.profile.models import Profile

PROFILE = Profile.model_validate(
    {
        "contact": {"name": "Jane Doe", "email": "jane@example.com", "phone": "555-0100", "location": "Austin, TX"},
        "summary": "Backend engineer with eight years of experience building billing systems for large customers.",
        "experience": [
            {
                "title": "Senior Engineer",
                "company": "Acme",
                "start": "January 2020",
                "current": True,
                "bullets": [
                    f"Built billing feature number {n} that cut invoice errors for twelve regional teams."
                    for n in range(12)
                ],
            }
        ],
        "skills": [{"category": "Languages", "items": ["C#", "Python", "SQL"]}],
    }
)

LONG = "Built a billing platform that reduced invoice errors for twelve regional finance teams."


def by_name(checks):
    return {c.name: c for c in checks}


def pdf_bytes(draw, *, font="Helvetica"):
    pdf = FPDF()
    pdf.set_margins(15, 15, 15)
    pdf.add_page()
    pdf.set_font(font, size=10)
    draw(pdf)
    return bytes(pdf.output())


def single_column(pdf):
    for n in range(30):
        pdf.multi_cell(0, 5, f"{LONG} Line {n}.", new_x="LMARGIN", new_y="NEXT")


def single_column_with_dates(pdf):
    for n in range(25):
        pdf.cell(130, 5, f"Senior Engineer at Company {n}")
        pdf.cell(0, 5, "Jan 2020 - Present", new_x="LMARGIN", new_y="NEXT", align="R")
        pdf.multi_cell(0, 5, LONG, new_x="LMARGIN", new_y="NEXT")


def two_columns(pdf):
    for n in range(24):
        y = pdf.get_y()
        pdf.set_xy(15, y)
        pdf.cell(55, 5, f"Skill {n}")
        pdf.set_xy(80, y)
        pdf.cell(0, 5, f"{LONG[:60]} number {n}", new_x="LMARGIN", new_y="NEXT")


def two_columns_word_by_word(pdf):
    """Word and Google Docs emit text in small pieces; draw each word separately to mimic that."""
    words = (LONG + " ").split()
    for n in range(26):
        y = pdf.get_y()
        for column_x in (15, 110):
            x = column_x
            for word in words[:9]:
                pdf.set_xy(x, y)
                pdf.cell(pdf.get_string_width(word + " "), 5, word)
                x += pdf.get_string_width(word + " ")
        pdf.set_y(y + 5)


def sidebar_with_main_column(pdf):
    """The common template: short skill items on the left, long paragraphs on the right."""
    for n in range(20):
        y = pdf.get_y()
        pdf.set_xy(15, y)
        pdf.cell(50, 5, ["Python", "C#", "SQL", "Azure", "Docker"][n % 5])
        pdf.set_xy(75, y)
        pdf.cell(0, 5, f"{LONG[:60]} number {n}", new_x="LMARGIN", new_y="NEXT")


def mostly_single_column_with_a_few_aligned_lines(pdf):
    for n in range(60):
        pdf.multi_cell(0, 5, f"{LONG} Line {n}.", new_x="LMARGIN", new_y="NEXT")
    for label in ("Languages", "Cloud", "Databases"):
        y = pdf.get_y()
        pdf.set_xy(15, y)
        pdf.cell(40, 5, label)
        pdf.set_xy(60, y)
        pdf.cell(0, 5, "C#, Python, SQL, TypeScript, Go, Rust, Java", new_x="LMARGIN", new_y="NEXT")


# --- PDF --------------------------------------------------------------------------------------------


def test_exported_style_pdf_passes_every_check():
    blocks = build_blocks(PROFILE)
    checks = by_name(check_pdf(build_pdf(blocks, name="Jane Doe", paper="letter").data))
    assert [c.name for c in checks.values() if not c.ok] == []


def test_single_column_with_right_aligned_dates_is_not_mistaken_for_columns():
    checks = by_name(check_pdf(pdf_bytes(single_column_with_dates)))
    assert checks["Single column"].ok, checks["Single column"].detail


def test_columns_drawn_word_by_word_are_still_detected():
    assert not by_name(check_pdf(pdf_bytes(two_columns_word_by_word)))["Single column"].ok


def test_a_template_sidebar_with_a_main_column_is_detected():
    assert not by_name(check_pdf(pdf_bytes(sidebar_with_main_column)))["Single column"].ok


def test_a_few_aligned_lines_in_a_mostly_single_column_resume_are_not_columns():
    checks = by_name(check_pdf(pdf_bytes(mostly_single_column_with_a_few_aligned_lines)))
    assert checks["Single column"].ok, checks["Single column"].detail


def test_plain_single_column_text_passes():
    assert by_name(check_pdf(pdf_bytes(single_column)))["Single column"].ok


def test_side_by_side_text_is_flagged_as_multi_column():
    checks = by_name(check_pdf(pdf_bytes(two_columns)))
    assert not checks["Single column"].ok
    assert "multi-column" in checks["Single column"].detail


def test_a_picture_of_a_resume_has_no_readable_text():
    checks = by_name(check_pdf(pdf_bytes(lambda pdf: None)))
    assert not checks["Text can be read"].ok
    assert "scan" in checks["Text can be read"].detail


def test_images_are_reported(tmp_path):
    from PIL import Image

    image = tmp_path / "logo.png"
    Image.new("RGB", (40, 40), "blue").save(image)

    def draw(pdf):
        single_column(pdf)
        pdf.image(str(image), x=160, y=10, w=20)

    checks = by_name(check_pdf(pdf_bytes(draw)))
    assert not checks["No images"].ok
    assert "ignore images" in checks["No images"].detail


def test_standard_fonts_are_accepted():
    assert by_name(check_pdf(pdf_bytes(single_column, font="Courier")))["Standard fonts"].ok
    assert by_name(check_pdf(pdf_bytes(single_column, font="Helvetica")))["Standard fonts"].ok


def test_garbled_text_encoding_is_detected(monkeypatch):
    import app.profile.file_checks as module

    real = module.PdfReader

    class Garbled(real):
        @property
        def pages(self):
            class Page:
                def __init__(self, inner):
                    self.inner = inner
                    self.mediabox = inner.mediabox
                    self.images = []

                def extract_text(self, *args, **kwargs):
                    return "(cid:12)(cid:7)�� " * 40

                def get(self, key, default=None):
                    return self.inner.get(key, default)

            return [Page(p) for p in super().pages]

    monkeypatch.setattr(module, "PdfReader", Garbled)
    checks = by_name(check_pdf(pdf_bytes(single_column)))
    assert not checks["Clean text encoding"].ok


def test_long_resumes_are_noted():
    def draw(pdf):
        for _ in range(3):
            single_column(pdf)
            pdf.add_page()

    assert not by_name(check_pdf(pdf_bytes(draw)))["Length"].ok


# --- DOCX -------------------------------------------------------------------------------------------


def test_exported_style_docx_passes_every_check():
    blocks = build_blocks(PROFILE)
    checks = by_name(check_docx(build_docx(blocks, name="Jane Doe", paper="letter")))
    assert [c.name for c in checks.values() if not c.ok] == []


def test_docx_with_a_table_header_and_unusual_font_is_flagged():
    source = make_docx(header="jane@example.com | 555-0100", table=[["Python", "Expert"]])
    document = docx.Document(io.BytesIO(source))
    document.paragraphs[0].runs[0].font.name = "Comic Sans MS"
    buffer = io.BytesIO()
    document.save(buffer)
    checks = by_name(check_docx(buffer.getvalue()))
    assert not checks["No tables"].ok
    assert not checks["No header or footer text"].ok
    assert not checks["Standard fonts"].ok
    assert "Comic Sans MS" in checks["Standard fonts"].detail


def test_docx_with_two_columns_is_flagged():
    document = docx.Document()
    document.add_paragraph("Jane Doe")
    cols = document.sections[0]._sectPr.xpath("./w:cols")[0]
    cols.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}num", "2")
    buffer = io.BytesIO()
    document.save(buffer)
    assert not by_name(check_docx(buffer.getvalue()))["Single column"].ok


# --- dispatch and import ----------------------------------------------------------------------------


def test_check_upload_dispatches_on_file_type_and_never_raises():
    assert check_upload("resume.pdf", pdf_bytes(single_column))
    assert check_upload("resume.docx", make_docx())
    assert check_upload("resume.txt", b"plain text") == []
    assert check_upload("broken.pdf", b"not a pdf") == []
    assert check_upload("broken.docx", b"not a docx") == []


@pytest.fixture
def configured(client, auth, store):
    store.set("sk-test")
    client.put("/settings", headers=auth, json={"model": "test/model"})
    return client


def test_importing_a_file_returns_checks_on_the_uploaded_file(configured, auth, monkeypatch):
    async def fake(**kwargs):
        return PROFILE.model_copy(deep=True)

    monkeypatch.setattr(llm, "complete_json", fake)
    docx_bytes = make_docx(table=[["Python", "Expert"]])
    result = configured.post("/profile/import", headers=auth, files={"file": ("resume.docx", docx_bytes)}).json()
    names = {c["name"]: c["ok"] for c in result["file_checks"]}
    assert names["No tables"] is False
    assert names["Single column"] is True


def test_pasted_text_has_no_file_checks(configured, auth, monkeypatch):
    async def fake(**kwargs):
        return PROFILE.model_copy(deep=True)

    monkeypatch.setattr(llm, "complete_json", fake)
    text = "Jane Doe, senior engineer with experience building billing systems for large customers."
    assert configured.post("/profile/import-text", headers=auth, json={"text": text}).json()["file_checks"] == []

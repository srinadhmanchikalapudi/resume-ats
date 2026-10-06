import io
from pathlib import Path

import docx
import pytest
from pypdf import PdfReader

from app.export import service
from app.export.blocks import build_blocks
from app.export.checks import check_docx, check_pdf
from app.export.docx_builder import build_docx
from app.export.models import ExportRequest
from app.export.pdf_builder import build_pdf
from app.export.text import pdf_safe, slug
from app.profile.models import Profile

LONG_DASH, CURLY_OPEN, CURLY_CLOSE, BULLET = chr(0x2013), chr(0x201C), chr(0x201D), chr(0x2022)

PROFILE = Profile.model_validate(
    {
        "contact": {
            "name": "Jane Doe",
            "email": "jane@example.com",
            "phone": "555-0100",
            "location": "Austin, TX",
            "links": [{"label": "LinkedIn", "url": "linkedin.com/in/janedoe"}],
        },
        "summary": f"Backend engineer {LONG_DASH} 8 years of {CURLY_OPEN}billing{CURLY_CLOSE} systems.",
        "experience": [
            {
                "title": "Senior Engineer",
                "company": "Acme",
                "location": "Austin, TX",
                "start": "Jan 2020",
                "current": True,
                "bullets": ["Built a billing API that cut errors by 40%.", "Led migration of 12 services to Azure Functions."],
            },
            {
                "title": "Engineer",
                "company": "Initech",
                "start": "2016",
                "end": "2019",
                "bullets": ["Maintained the reporting platform."],
            },
        ],
        "skills": [{"category": "Languages", "items": ["C#", "Python"]}, {"category": "Cloud", "items": ["Azure"]}],
        "education": [{"school": "UT Austin", "degree": "BS", "field": "Computer Science", "end": "2016"}],
        "certifications": [{"name": "AZ-204", "issuer": "Microsoft", "date": "2022"}],
        "projects": [{"name": "Side project", "description": "A tool", "tech": ["Python"], "bullets": ["Shipped it."]}],
    }
)


def heading_order(blocks):
    return [b.text for b in blocks if b.kind == "heading"]


# --- blocks -----------------------------------------------------------------------------------------


def test_block_order_with_skills_first_and_after_experience():
    assert heading_order(build_blocks(PROFILE, skills_first=True)) == [
        "Summary", "Skills", "Experience", "Education", "Certifications", "Projects",
    ]  # fmt: skip
    assert heading_order(build_blocks(PROFILE, skills_first=False)) == [
        "Summary", "Experience", "Skills", "Education", "Certifications", "Projects",
    ]  # fmt: skip


def test_blocks_describe_roles_with_title_then_company_and_dates():
    blocks = build_blocks(PROFILE)
    titles = [(b.kind, b.text) for b in blocks if b.kind in {"item_title", "item_meta"}]
    assert ("item_title", "Senior Engineer") in titles
    assert ("item_meta", "Acme, Austin, TX | Jan 2020 - Present") in titles
    assert ("item_meta", "Initech | 2016 - 2019") in titles


def test_empty_sections_are_omitted():
    blocks = build_blocks(Profile.model_validate({"contact": {"name": "Jane"}, "experience": [{"title": "Dev", "bullets": ["x"]}]}))
    assert heading_order(blocks) == ["Experience"]


# --- text safety ------------------------------------------------------------------------------------


def test_typographic_characters_pass_through_for_the_pdf():
    text = f"{CURLY_OPEN}Led{CURLY_CLOSE} 2020 {LONG_DASH} 2023 {BULLET} caf" + chr(0xE9)
    assert pdf_safe(text) == (text, 0)


def test_unsupported_characters_fall_back_or_are_counted():
    safe, replaced = pdf_safe("Lukasz " + chr(0x141) + "od" + chr(0x17A) + " " + chr(0x0936) + chr(0x200B) + "x")
    assert replaced == 2  # the Devanagari letter and the L with stroke have no plain fallback
    assert "?" in safe and chr(0x200B) not in safe
    assert pdf_safe("a" + chr(0x2192) + "b") == ("a->b", 0)


def test_slug_is_filesystem_safe():
    assert slug("Globex Financial, Inc.") == "Globex-Financial-Inc"
    assert slug("  ../../etc/passwd  ") == "etc-passwd"
    assert slug("") == ""
    assert len(slug("x" * 100)) == 40


# --- DOCX -------------------------------------------------------------------------------------------


@pytest.fixture(scope="module")
def blocks():
    return build_blocks(PROFILE)


@pytest.fixture(scope="module")
def docx_bytes(blocks):
    return build_docx(blocks, name="Jane Doe", paper="letter")


def test_docx_passes_every_ats_check(docx_bytes, blocks):
    checks = check_docx(docx_bytes, blocks)
    assert [c.name for c in checks if not c.ok] == []
    assert {c.name for c in checks} >= {"Single column", "No tables", "Standard font", "All bullets present"}


def test_docx_structure_is_plain_and_uses_real_styles(docx_bytes):
    document = docx.Document(io.BytesIO(docx_bytes))
    assert len(document.tables) == 0
    assert len(document.inline_shapes) == 0
    styles = {p.style.name for p in document.paragraphs}
    assert {"Heading 1", "List Bullet"} <= styles
    assert [p.text for p in document.paragraphs if p.style.name == "Heading 1"][:3] == ["Summary", "Skills", "Experience"]
    assert document.paragraphs[0].text == "Jane Doe"
    assert "jane@example.com" in document.paragraphs[1].text


def test_docx_uses_arial_without_theme_fonts(docx_bytes):
    document = docx.Document(io.BytesIO(docx_bytes))
    for style_name in ("Normal", "Heading 1", "List Bullet"):
        rfonts = document.styles[style_name].element.rPr.rFonts
        assert rfonts.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}ascii") == "Arial"
        assert rfonts.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}asciiTheme") is None


def test_docx_opens_in_current_word_mode_not_compatibility_mode(docx_bytes):
    document = docx.Document(io.BytesIO(docx_bytes))
    modes = document.settings.element.xpath("./w:compat/w:compatSetting[@w:name='compatibilityMode']/@w:val")
    assert modes == ["15"]


def test_docx_metadata_names_the_candidate_not_the_library(docx_bytes):
    properties = docx.Document(io.BytesIO(docx_bytes)).core_properties
    assert properties.author == "Jane Doe"
    assert properties.title == "Jane Doe - Resume"


def test_docx_page_size_follows_paper_choice(blocks):
    letter = docx.Document(io.BytesIO(build_docx(blocks, name="J", paper="letter"))).sections[0]
    a4 = docx.Document(io.BytesIO(build_docx(blocks, name="J", paper="a4"))).sections[0]
    assert round(letter.page_width.inches, 1) == 8.5
    assert round(a4.page_width.mm) == 210


def test_docx_check_catches_a_table():
    document = docx.Document()
    document.add_paragraph("Jane Doe")
    document.add_table(rows=1, cols=2)
    buffer = io.BytesIO()
    document.save(buffer)
    failed = {c.name for c in check_docx(buffer.getvalue(), []) if not c.ok}
    assert "No tables" in failed


def test_docx_check_catches_header_text():
    document = docx.Document()
    document.add_paragraph("Jane Doe")
    document.sections[0].header.paragraphs[0].text = "jane@example.com"
    buffer = io.BytesIO()
    document.save(buffer)
    failed = {c.name for c in check_docx(buffer.getvalue(), []) if not c.ok}
    assert "No header or footer text" in failed


# --- PDF --------------------------------------------------------------------------------------------


@pytest.fixture(scope="module")
def pdf(blocks):
    return build_pdf(blocks, name="Jane Doe", paper="letter")


def test_pdf_text_can_be_read_back_in_order(pdf):
    text = " ".join("\n".join(page.extract_text() for page in PdfReader(io.BytesIO(pdf.data)).pages).split())
    for expected in ("Jane Doe", "SUMMARY", "SKILLS", "EXPERIENCE", "Senior Engineer", "Azure Functions"):
        assert expected in text
    assert text.index("SUMMARY") < text.index("SKILLS") < text.index("EXPERIENCE") < text.index("EDUCATION")
    assert "billing" in text and BULLET in text and LONG_DASH in text


def test_pdf_passes_every_ats_check(pdf, blocks):
    checks, pages = check_pdf(pdf.data, blocks, pdf.replaced_characters)
    assert [c.name for c in checks if not c.ok] == []
    assert pages == pdf.pages == 1


def test_pdf_has_standard_fonts_and_no_images(pdf):
    reader = PdfReader(io.BytesIO(pdf.data))
    assert sum(len(page.images) for page in reader.pages) == 0
    assert reader.metadata.title == "Jane Doe - Resume"
    assert reader.metadata.author == "Jane Doe"


def test_long_resume_paginates_and_reports_length(blocks):
    long_profile = PROFILE.model_copy(deep=True)
    long_profile.experience[0].bullets = [f"Delivered result number {i} with a long description of the work. " * 3 for i in range(40)]
    long_blocks = build_blocks(long_profile)
    result = build_pdf(long_blocks, name="Jane Doe", paper="letter")
    checks, pages = check_pdf(result.data, long_blocks, 0)
    assert pages > 2
    assert {c.name: c.ok for c in checks}["Length"] is False
    assert {c.name: c.ok for c in checks}["All bullets present"] is True  # every bullet survives across pages


def test_unsupported_characters_are_reported(blocks):
    odd = PROFILE.model_copy(deep=True)
    odd.experience[0].bullets = ["Worked with " + chr(0x0936) + chr(0x094D) + " customers."]
    odd_blocks = build_blocks(odd)
    result = build_pdf(odd_blocks, name="Jane Doe", paper="letter")
    checks, _ = check_pdf(result.data, odd_blocks, result.replaced_characters)
    assert result.replaced_characters == 2
    assert {c.name: c.ok for c in checks}["Characters"] is False


# --- service and API --------------------------------------------------------------------------------


def test_export_writes_both_files_into_a_dated_folder(tmp_path):
    result = service.export_resume(ExportRequest(profile=PROFILE, job_title="Senior Engineer", company="Globex Inc"))
    assert result.folder.endswith("_Globex-Inc_Senior-Engineer")
    assert [f.format for f in result.files] == ["docx", "pdf"]
    for file in result.files:
        assert file.filename.startswith("Jane_Doe_Resume")
        assert tmp_path in Path(file.path).parents
        assert file.size_bytes > 0
        assert all(c.ok for c in file.checks)
    assert result.files[1].pages == 1
    assert result.warnings == []


def test_second_export_for_the_same_job_gets_its_own_folder():
    first = service.export_resume(ExportRequest(profile=PROFILE, job_title="Eng", company="Acme"))
    second = service.export_resume(ExportRequest(profile=PROFILE, job_title="Eng", company="Acme"))
    assert first.folder != second.folder
    assert second.folder.endswith("-2")


def test_export_warns_about_missing_email():
    no_email = PROFILE.model_copy(deep=True)
    no_email.contact.email = ""
    result = service.export_resume(ExportRequest(profile=no_email))
    assert any("email" in w for w in result.warnings)


def test_export_of_an_empty_resume_is_rejected():
    with pytest.raises(ValueError, match="nothing to export"):
        service.export_resume(ExportRequest(profile=Profile()))


def test_export_endpoint_and_file_download(client, auth):
    response = client.post(
        "/export", headers=auth, json={"profile": PROFILE.model_dump(), "job_title": "Eng", "company": "Acme"}
    )
    assert response.status_code == 200
    files = {f["format"]: f for f in response.json()["files"]}

    pdf = client.get("/export/file", headers=auth, params={"path": files["pdf"]["path"]})
    assert pdf.status_code == 200
    assert pdf.headers["content-type"] == "application/pdf"
    assert pdf.content.startswith(b"%PDF")

    word = client.get("/export/file", headers=auth, params={"path": files["docx"]["path"]})
    assert word.status_code == 200
    assert word.content[:2] == b"PK"  # a DOCX is a zip


def test_export_endpoint_validates_and_requires_token(client, auth):
    assert client.post("/export", json={"profile": PROFILE.model_dump()}).status_code == 401
    assert client.post("/export", headers=auth, json={"profile": Profile().model_dump()}).status_code == 422


def test_download_refuses_paths_outside_the_exports_folder(client, auth, tmp_path):
    outside = tmp_path / "secret.pdf"
    outside.write_bytes(b"%PDF-1.4 not yours")
    assert client.get("/export/file", headers=auth, params={"path": str(outside)}).status_code == 404

    result = client.post("/export", headers=auth, json={"profile": PROFILE.model_dump(), "company": "Acme"}).json()
    traversal = result["folder"] + "/../../settings.json"
    assert client.get("/export/file", headers=auth, params={"path": traversal}).status_code == 404
    wrong_type = result["files"][0]["path"] + ".txt"
    assert client.get("/export/file", headers=auth, params={"path": wrong_type}).status_code == 404

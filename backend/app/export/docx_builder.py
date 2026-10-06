"""ATS-safe DOCX: one column, real heading and list styles, no tables, images, text boxes, headers or footers."""

from __future__ import annotations

import io

import docx
from docx.enum.text import WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt, RGBColor

from .blocks import Block
from .models import Paper

FONT = "Arial"
BLACK = RGBColor(0, 0, 0)
GRAY = RGBColor(0x40, 0x40, 0x40)

_PPR_AFTER_BORDER = (
    "w:shd", "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap", "w:overflowPunct", "w:topLinePunct",
    "w:autoSpaceDE", "w:autoSpaceDN", "w:bidi", "w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind",
    "w:contextualSpacing", "w:mirrorIndents", "w:suppressOverlap", "w:jc", "w:textDirection", "w:textAlignment",
    "w:textboxTightWrap", "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange",
)  # fmt: skip


def _set_font(style, size: float, *, bold: bool | None = None, color: RGBColor = BLACK) -> None:
    """Forces a plain font. Theme font attributes are removed because Word lets them override the name."""
    style.font.name = FONT
    style.font.size = Pt(size)
    style.font.color.rgb = color
    if bold is not None:
        style.font.bold = bold
    rfonts = style.element.get_or_add_rPr().find(qn("w:rFonts"))
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), FONT)
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if rfonts.get(qn(attr)) is not None:
            del rfonts.attrib[qn(attr)]


def _bottom_rule(style) -> None:
    border = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for key, value in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", "808080")):
        bottom.set(qn(key), value)
    border.append(bottom)
    style.element.get_or_add_pPr().insert_element_before(border, *_PPR_AFTER_BORDER)


def _use_current_word_mode(document) -> None:
    """The template says Word 2010 (14), which makes Word print "[Compatibility Mode]" in the title bar."""
    for setting in document.settings.element.xpath("./w:compat/w:compatSetting[@w:name='compatibilityMode']"):
        setting.set(qn("w:val"), "15")


def build_docx(blocks: list[Block], *, name: str, paper: Paper) -> bytes:
    document = docx.Document()
    _use_current_word_mode(document)

    section = document.sections[0]
    if paper == "a4":
        section.page_width, section.page_height = Mm(210), Mm(297)
    else:
        section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.left_margin = section.right_margin = Inches(0.75)
    section.top_margin = section.bottom_margin = Inches(0.7)

    styles = document.styles
    normal = styles["Normal"]
    _set_font(normal, 10)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(2)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE

    heading = styles["Heading 1"]
    _set_font(heading, 11, bold=True)
    heading.paragraph_format.space_before = Pt(10)
    heading.paragraph_format.space_after = Pt(4)
    heading.paragraph_format.keep_with_next = True
    _bottom_rule(heading)

    bullet_style = styles["List Bullet"]
    _set_font(bullet_style, 10)
    bullet_style.paragraph_format.left_indent = Inches(0.25)
    bullet_style.paragraph_format.first_line_indent = Inches(-0.18)
    bullet_style.paragraph_format.space_after = Pt(2)

    for block in blocks:
        if block.kind == "heading":
            document.add_paragraph(block.text, style="Heading 1")
        elif block.kind == "bullet":
            document.add_paragraph(block.text, style="List Bullet")
        else:
            paragraph = document.add_paragraph()
            run = paragraph.add_run(block.text)
            if block.kind == "name":
                run.bold = True
                run.font.size = Pt(18)
                paragraph.paragraph_format.space_after = Pt(3)
            elif block.kind == "contact":
                run.font.size = Pt(9.5)
                run.font.color.rgb = GRAY
            elif block.kind == "item_title":
                run.bold = True
                paragraph.paragraph_format.space_before = Pt(6)
                paragraph.paragraph_format.keep_with_next = True
            elif block.kind == "item_meta":
                run.font.size = Pt(9.5)
                run.font.color.rgb = GRAY
                paragraph.paragraph_format.keep_with_next = True

    properties = document.core_properties
    properties.author = name
    properties.last_modified_by = name
    properties.title = f"{name} - Resume" if name else "Resume"
    properties.comments = ""
    properties.subject = ""

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()

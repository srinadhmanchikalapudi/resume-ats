"""Verifies the generated files by reading them back the way an applicant tracking system would."""

from __future__ import annotations

import io

import docx
from pypdf import PdfReader

from .blocks import SECTION_TITLES, Block
from .models import Check
from .text import pdf_safe

ATS_FONTS = {
    "arial", "calibri", "helvetica", "times new roman", "times", "times-roman", "georgia", "cambria",
    "garamond", "verdana", "tahoma", "courier",
}  # fmt: skip
MAX_PAGES = 2


def _norm(text: str) -> str:
    return " ".join(text.split())


def _bullets(blocks: list[Block]) -> list[str]:
    return [block.text for block in blocks if block.kind == "bullet"]


def check_docx(data: bytes, blocks: list[Block]) -> list[Check]:
    document = docx.Document(io.BytesIO(data))
    body_xml = document.element.xml
    paragraphs = [p for p in document.paragraphs]
    checks: list[Check] = []

    columns = [int(n) for section in document.sections for n in section._sectPr.xpath("./w:cols/@w:num")]
    checks.append(
        Check(name="Single column", ok=all(n <= 1 for n in columns), detail="One text column, no side panels.")
    )

    checks.append(Check(name="No tables", ok=len(document.tables) == 0, detail="Content is plain paragraphs."))

    has_graphics = bool(document.inline_shapes) or any(
        marker in body_xml for marker in ("<w:drawing", "<w:pict", "<w:txbxContent")
    )
    checks.append(Check(name="No images or text boxes", ok=not has_graphics))

    has_header_footer = any(
        not part.is_linked_to_previous and any(p.text.strip() for p in part.paragraphs)
        for section in document.sections
        for part in (section.header, section.footer)
    )
    checks.append(
        Check(
            name="No header or footer text",
            ok=not has_header_footer,
            detail="Many systems skip headers and footers, so all details are in the body.",
        )
    )

    fonts = {document.styles["Normal"].font.name, document.styles["Heading 1"].font.name}
    fonts |= {run.font.name for p in paragraphs for run in p.runs if run.font.name}
    bad_fonts = sorted(f for f in fonts if f and f.lower() not in ATS_FONTS)
    checks.append(
        Check(name="Standard font", ok=not bad_fonts, detail=", ".join(sorted(f for f in fonts if f)) or "default")
    )

    headings = [p.text.strip() for p in paragraphs if p.style.name == "Heading 1"]
    unknown = [h for h in headings if h.lower() not in {t.lower() for t in SECTION_TITLES}]
    checks.append(
        Check(
            name="Standard section headings",
            ok=bool(headings) and not unknown,
            detail=", ".join(headings) if headings else "No section headings found.",
        )
    )

    contact = next((b.text for b in blocks if b.kind == "contact"), "")
    top = [_norm(p.text) for p in paragraphs[:4]]
    checks.append(
        Check(
            name="Contact details in the body",
            ok=bool(contact) and _norm(contact) in top,
            detail="Placed at the top of the page." if contact else "No contact details on the resume.",
        )
    )

    present = {_norm(p.text) for p in paragraphs}
    bullets = _bullets(blocks)
    found = sum(_norm(b) in present for b in bullets)
    checks.append(
        Check(name="All bullets present", ok=found == len(bullets), detail=f"{found} of {len(bullets)} bullets.")
    )
    return checks


def _pdf_fonts(reader: PdfReader) -> set[str]:
    names: set[str] = set()
    for page in reader.pages:
        resources = page.get("/Resources")
        fonts = resources.get_object().get("/Font") if resources else None
        if fonts is None:
            continue
        for font in fonts.get_object().values():
            base = str(font.get_object().get("/BaseFont", "")).lstrip("/")
            if base:
                names.add(base)
    return names


def check_pdf(data: bytes, blocks: list[Block], replaced: int) -> tuple[list[Check], int]:
    reader = PdfReader(io.BytesIO(data))
    pages = len(reader.pages)
    text = _norm("\n".join(page.extract_text() or "" for page in reader.pages))
    checks: list[Check] = []

    checks.append(
        Check(name="Text layer", ok=len(text) > 0, detail=f"{len(text):,} characters can be read back from the file.")
    )

    cursor, in_order = 0, True
    for block in blocks:
        if block.kind in {"name", "heading"}:
            expected = _norm(pdf_safe(block.text.upper() if block.kind == "heading" else block.text)[0])
            found = text.find(expected, cursor)
            if found == -1:
                in_order = False
                break
            cursor = found
    checks.append(
        Check(
            name="Reading order",
            ok=in_order,
            detail="Name and section headings come back in the intended order.",
        )
    )

    bullets = _bullets(blocks)
    found_count = sum(_norm(pdf_safe(b)[0]) in text for b in bullets)
    checks.append(
        Check(name="All bullets present", ok=found_count == len(bullets), detail=f"{found_count} of {len(bullets)} bullets.")
    )

    images = sum(len(page.images) for page in reader.pages)
    checks.append(Check(name="No images", ok=images == 0))

    fonts = _pdf_fonts(reader)
    bad = sorted(f for f in fonts if f.split("-")[0].lower() not in ATS_FONTS)
    checks.append(Check(name="Standard font", ok=not bad, detail=", ".join(sorted(fonts)) or "none"))

    checks.append(
        Check(
            name="Characters",
            ok=replaced == 0,
            detail="All characters are shown."
            if replaced == 0
            else f"{replaced} character(s) the PDF fonts cannot draw became '?'. Use the DOCX for those.",
        )
    )

    checks.append(
        Check(
            name="Length",
            ok=pages <= MAX_PAGES,
            detail=f"{pages} page{'s' if pages != 1 else ''}."
            + ("" if pages <= MAX_PAGES else " Most recruiters read two at most; try the Concise length."),
        )
    )
    return checks, pages

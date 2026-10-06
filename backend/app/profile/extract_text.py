"""Plain-text extraction from uploaded resumes (PDF, DOCX, TXT/MD)."""

from __future__ import annotations

import io
from dataclasses import dataclass, field
from pathlib import PurePath

import docx
from pypdf import PdfReader
from pypdf.errors import PyPdfError

MIN_TEXT_CHARS = 40
THIN_TEXT_CHARS = 300


class ExtractError(ValueError):
    """The file cannot be turned into usable text; the message is safe to show to the user."""


@dataclass
class ExtractedText:
    text: str
    warnings: list[str] = field(default_factory=list)


def extract_text(filename: str, data: bytes) -> ExtractedText:
    suffix = PurePath(filename).suffix.lower()
    if suffix == ".pdf":
        result = _from_pdf(data)
    elif suffix == ".docx":
        result = _from_docx(data)
    elif suffix in {".txt", ".md"}:
        result = ExtractedText(data.decode("utf-8", errors="replace"))
    elif suffix == ".doc":
        raise ExtractError("Old .doc files are not supported. Save it as .docx or PDF and try again.")
    else:
        raise ExtractError("Unsupported file type. Use a PDF, DOCX or plain-text file.")

    result.text = result.text.strip()
    if len(result.text) < MIN_TEXT_CHARS:
        raise ExtractError(
            "Almost no text could be read from this file. If it is a scanned or image-only PDF, "
            "it also cannot be read by most applicant tracking systems; export a text-based PDF "
            "or use a DOCX instead."
        )
    if len(result.text) < THIN_TEXT_CHARS:
        result.warnings.append("Very little text was found. Check that the whole resume was read.")
    return result


def _from_pdf(data: bytes) -> ExtractedText:
    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted and not reader.decrypt(""):
            raise ExtractError("This PDF is password protected. Remove the password and try again.")
        pages = [page.extract_text() or "" for page in reader.pages]
    except ExtractError:
        raise
    except (PyPdfError, ValueError, OSError) as exc:
        raise ExtractError(f"This PDF could not be read ({exc}).") from exc
    return ExtractedText("\n\n".join(pages))


def _from_docx(data: bytes) -> ExtractedText:
    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:  # python-docx raises several unrelated types for bad zips/XML
        raise ExtractError(f"This DOCX could not be read ({exc}).") from exc

    warnings: list[str] = []
    parts: list[str] = []

    header_footer: list[str] = []
    for section in document.sections:
        for container in (section.header, section.footer):
            header_footer.extend(p.text.strip() for p in container.paragraphs if p.text.strip())
    if header_footer:
        warnings.append(
            "Text was found in the document header or footer. Many applicant tracking systems "
            "skip headers and footers, so keep contact details in the main body."
        )
        parts.extend(dict.fromkeys(header_footer))  # de-duplicate repeated per-section headers

    parts.extend(p.text for p in document.paragraphs)

    if document.tables:
        warnings.append(
            "Tables were found. Some applicant tracking systems read table content in the wrong "
            "order or skip it; a plain single-column layout is safer."
        )
        for table in document.tables:
            for row in table.rows:
                cells = dict.fromkeys(cell.text.strip() for cell in row.cells if cell.text.strip())
                if cells:
                    parts.append(" | ".join(cells))

    return ExtractedText("\n".join(parts), warnings)
